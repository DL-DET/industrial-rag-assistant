from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from file_parsers import (
    decode_text_file,
    extract_docx_text,
    extract_pdf_text,
    extract_text_from_upload,
)


def test_decode_text_file_supports_gb18030():
    assert decode_text_file("工业风机".encode("gb18030")) == "工业风机"


def test_extract_text_from_txt_upload():
    assert extract_text_from_upload("说明.txt", "设备说明".encode()) == "设备说明"


def test_extract_text_from_docx_upload():
    document = Document()
    document.add_paragraph("离心风机维护说明")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "参数"
    table.cell(0, 1).text = "额定转速"
    buffer = BytesIO()
    document.save(buffer)

    text = extract_docx_text(buffer.getvalue())

    assert "离心风机维护说明" in text
    assert "参数 | 额定转速" in text


def test_unsupported_file_type_raises_error():
    with pytest.raises(ValueError, match="不支持的文件类型"):
        extract_text_from_upload("设备.csv", b"a,b")


def test_empty_text_file_raises_error():
    with pytest.raises(ValueError, match="未从文件中提取到文本"):
        extract_text_from_upload("空文件.txt", b"   ")


def test_blank_pdf_does_not_fail_on_none_text():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = BytesIO()
    writer.write(buffer)

    assert extract_pdf_text(buffer.getvalue()) == ""
